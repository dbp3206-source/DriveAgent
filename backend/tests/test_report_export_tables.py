import io

import pdfplumber
from docx import Document

from app.services.report_exports import _blocks, _pdf_text, export_docx, export_pdf

CONTENT = "## Dữ kiện\n\n| Nội dung | Giá trị |\n| --- | ---: |\n| Nhân viên | 42 |\n"


def test_word_export_has_native_table_and_vietnamese(tmp_path):
    body = export_docx("Chuẩn bị tư vấn", CONTENT)
    (tmp_path / "consultation.docx").write_bytes(body)
    document = Document(io.BytesIO(body))
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "Nhân viên"
    assert document.tables[0].cell(1, 1).text == "42"
    assert document.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader")


def test_pdf_table_export_is_valid(tmp_path):
    result = export_pdf("Chuẩn bị tư vấn", CONTENT +
                        "\n[Nguồn chính](https://example.com/source)\n")
    (tmp_path / "consultation.pdf").write_bytes(result)
    assert result.startswith(b"%PDF-")
    assert len(result) > 1000


def test_long_pdf_table_keeps_rows_and_repeats_heading():
    table = "| Mục | Giá trị |\n| --- | --- |\n" + "\n".join(
        f"| Item-{index:03d} | Value-{index:03d} |" for index in range(90)
    )
    with pdfplumber.open(io.BytesIO(export_pdf("Bảng dài kiểm thử", table))) as reader:
        assert len(reader.pages) >= 3
        pages = [page.extract_text() for page in reader.pages]
    assert all("Giá trị" in page for page in pages)
    text = "\n".join(pages)
    assert all(f"Item-{index:03d}" in text and f"Value-{index:03d}" in text
               for index in range(90))


def test_invalid_table_does_not_discard_text():
    assert list(_blocks("A | B\n--- | not a separator")) == [
        ("line", "A | B"), ("line", "--- | not a separator"),
    ]


def test_escaped_pipe_stays_inside_cell():
    assert list(_blocks(r"| A\|B | C |" + "\n| --- | --- |"))[0] == (
        "table", [["A|B", "C"]],
    )


def test_word_source_link_is_real_and_bold_is_not_raw_markdown():
    document = Document(io.BytesIO(export_docx(
        "Nguồn", "**Đã xác nhận** [Nguồn chính](https://example.com/source)",
    )))
    assert any(run.bold and run.text == "Đã xác nhận" for p in document.paragraphs
               for run in p.runs)
    assert any(rel.target_ref == "https://example.com/source"
               for rel in document.part.rels.values())
    assert document.paragraphs[-1]._p.xpath("./w:hyperlink/w:r/w:t")[0].text == "Nguồn chính"


def test_pdf_inline_links_escape_untrusted_markup_and_reject_other_protocols():
    assert _pdf_text('<img src="secret">') == '&lt;img src="secret"&gt;'
    assert '<link ' not in _pdf_text('[Không an toàn](javascript:alert)')
    assert 'href="https://example.com?a=1&amp;b=2"' in _pdf_text(
        '[Nguồn](https://example.com?a=1&b=2)',
    )
